import io
from openpyxl import load_workbook
from kalkulasjon import ProjectInput,beregn_prosjekt
from export_excel import lag_excel
from fdv import lag_fdv_oversikt
shapes=[[(0,0),(4,0),(4,3),(0,3)],[(0,0),(6,0),(6,3),(3,3),(3,5),(0,5)],[(0,0),(6,0),(6,3),(5,4),(0,4)]]
count=0
for pts in shapes:
 for kind in ('Kjølerom','Fryserom'):
  for thickness in (80,100,150):
   panel=f'PIR {thickness} mm'
   inp=ProjectInput('Excelkontroll','138','A','22.10.2026','1:50',kind,panel,pts,2.4,1.1,[],[])
   result=beregn_prosjekt(inp)
   book=load_workbook(io.BytesIO(lag_excel(result,lag_fdv_oversikt(panel,kind,[],[]))))
   assert set(book.sheetnames)=={'Materialliste','Kappliste','Sammendrag','FDV'}
   summary={row[0]:row[1] for row in list(book['Sammendrag'].values)[1:]}
   assert abs(summary['Priset delsum']-result['tot_pris'])<1e-7, (summary,result['tot_pris'])
   assert summary['Manglende priser']
   assert book['Materialliste'].max_row==len(result['df_mat'])+1
   assert book['Kappliste'].max_row==len(result['df_kapp'])+1
   assert any('PE25' in str(c) for row in book['FDV'].values for c in row)
   count+=1
print(f'PASS: {count} Excel exports match calculation/cuts, mark incomplete price and include PE25 FDV')
